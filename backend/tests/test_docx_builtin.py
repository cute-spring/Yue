import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from pydantic_ai import RunContext

from app.mcp.builtin.docx import (
    DocxExtractTablesTool,
    DocxChangesTool,
    DocxCommentsTool,
    DocxLinksTool,
    DocxMediaTool,
    DocxMetadataTool,
    DocxNotesTool,
    DocxProfileTool,
    DocxQueryTool,
    DocxReadTool,
    DocxSearchTool,
    DocxStructureTool,
    DocxRenderTool,
)
from app.mcp.builtin.registry import builtin_tool_registry


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    path = tmp_path / "brief.docx"
    document = Document()
    document.core_properties.title = "Product brief"
    document.core_properties.author = "Yue"
    document.add_heading("Summary", level=1)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Ready"
    document.save(path)
    return path


@pytest.fixture
def mock_ctx():
    return MagicMock(spec=RunContext)


@pytest.mark.asyncio
async def test_docx_profile_returns_document_summary(mock_ctx, sample_docx):
    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(sample_docx.parent)], [])):
        payload = json.loads(
            await DocxProfileTool().execute(mock_ctx, {"path": sample_docx.name, "root_dir": str(sample_docx.parent)})
        )

    assert payload["ok"] is True
    assert payload["tool"] == "docx_profile"
    assert payload["metadata"]["title"] == "Product brief"
    assert payload["counts"] == {"headings": 1, "paragraphs": 3, "tables": 1, "sections": 1}
    assert payload["heading_outline"] == [{"text": "Summary", "level": 1, "block_id": "p-0001"}]
    assert payload["features"]["has_tables"] is True


@pytest.mark.asyncio
async def test_docx_read_paginates_blocks_and_renders_markdown(mock_ctx, sample_docx):
    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(sample_docx.parent)], [])):
        payload = json.loads(
            await DocxReadTool().execute(
                mock_ctx,
                {"path": sample_docx.name, "root_dir": str(sample_docx.parent), "limit": 2, "mode": "markdown"},
            )
        )

    assert payload["ok"] is True
    assert payload["tool"] == "docx_read"
    assert payload["is_truncated"] is True
    assert payload["next_cursor"] == 2
    assert "# Summary" in payload["data"]
    assert payload["citations"][0]["block_id"] == "p-0001"


@pytest.mark.asyncio
async def test_docx_extract_tables_selects_a_table_by_id(mock_ctx, sample_docx):
    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(sample_docx.parent)], [])):
        payload = json.loads(
            await DocxExtractTablesTool().execute(
                mock_ctx,
                {"path": sample_docx.name, "root_dir": str(sample_docx.parent), "table_index": 1},
            )
        )

    assert payload["ok"] is True
    assert payload["tool"] == "docx_extract_tables"
    assert payload["tables"][0]["table_id"] == "table-0004"
    assert payload["tables"][0]["rows"][1][1]["text"] == "Ready"


@pytest.mark.asyncio
async def test_docx_extract_tables_bounds_large_table_cells(mock_ctx, tmp_path: Path):
    path = tmp_path / "large-table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=501)
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = str(index)
    document.save(path)

    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(tmp_path)], [])):
        payload = json.loads(
            await DocxExtractTablesTool().execute(
                mock_ctx,
                {"path": path.name, "root_dir": str(tmp_path)},
            )
        )

    assert payload["ok"] is True
    assert payload["is_truncated"] is True
    assert len(payload["tables"][0]["rows"][0]) == 500


def test_registry_contains_docx_reader_tools():
    names = [tool.name for tool in builtin_tool_registry.get_all_tools()]
    assert {"docx_profile", "docx_read", "docx_extract_tables"}.issubset(names)


@pytest.mark.asyncio
async def test_docx_retrieval_tools_return_cited_and_structured_results(mock_ctx, sample_docx):
    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(sample_docx.parent)], [])):
        search = json.loads(
            await DocxSearchTool().execute(
                mock_ctx,
                {"path": sample_docx.name, "root_dir": str(sample_docx.parent), "query": "ready", "domains": ["table_cells"]},
            )
        )
        query = json.loads(
            await DocxQueryTool().execute(
                mock_ctx,
                {"path": sample_docx.name, "root_dir": str(sample_docx.parent), "query": {"kind": "paragraph"}},
            )
        )
        structure = json.loads(
            await DocxStructureTool().execute(
                mock_ctx,
                {"path": sample_docx.name, "root_dir": str(sample_docx.parent)},
            )
        )

    assert search["tool"] == "docx_search"
    assert search["matches"][0]["citation"]["kind"] == "table_cell"
    assert query["tool"] == "docx_query"
    assert query["applied_filters"] == {"kind": "paragraph"}
    assert structure["tool"] == "docx_structure"
    assert structure["headings"][0]["text"] == "Summary"


def test_registry_contains_docx_retrieval_tools():
    names = [tool.name for tool in builtin_tool_registry.get_all_tools()]
    assert {"docx_search", "docx_query", "docx_structure"}.issubset(names)


@pytest.mark.asyncio
async def test_docx_metadata_and_render_tools_register_schema_and_return_structured_results(mock_ctx, sample_docx, monkeypatch):
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.metadata", lambda *_: {"ok": True, "core": {}})
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.render", lambda *_: {"ok": False, "error_code": "DOCX_RENDERER_UNAVAILABLE", "message": "Install LibreOffice."})
    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(sample_docx.parent)], [])):
        metadata = json.loads(await DocxMetadataTool().execute(mock_ctx, {"path": sample_docx.name, "root_dir": str(sample_docx.parent)}))
        rendered = json.loads(await DocxRenderTool().execute(mock_ctx, {"path": sample_docx.name, "page_start": 1, "page_end": 2, "root_dir": str(sample_docx.parent)}))

    assert metadata == {"ok": True, "core": {}, "tool": "docx_metadata"}
    assert rendered == {"ok": False, "error_code": "DOCX_RENDERER_UNAVAILABLE", "message": "Install LibreOffice.", "tool": "docx_render"}
    assert DocxRenderTool().parameters["properties"]["page_end"]["minimum"] == 1
    names = [tool.name for tool in builtin_tool_registry.get_all_tools()]
    assert {"docx_metadata", "docx_render"}.issubset(names)


@pytest.mark.asyncio
async def test_docx_review_content_tools_register_bound_schemas_and_execute(mock_ctx, sample_docx, monkeypatch):
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.comments", lambda *_args, **_kwargs: {"ok": True, "comments": []})
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.changes", lambda *_args, **_kwargs: {"ok": True, "changes": []})
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.notes", lambda *_args, **_kwargs: {"ok": True, "notes": []})
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.media", lambda *_args, **_kwargs: {"ok": True, "items": []})
    monkeypatch.setattr("app.mcp.builtin.docx.docx_service.links", lambda *_args, **_kwargs: {"ok": True, "links": []})
    with patch("app.mcp.builtin.docx._get_doc_access", return_value=([str(sample_docx.parent)], [])):
        payloads = [
            json.loads(await tool.execute(mock_ctx, {"path": sample_docx.name, "limit": 1, "root_dir": str(sample_docx.parent)}))
            for tool in (DocxCommentsTool(), DocxChangesTool(), DocxNotesTool(), DocxMediaTool(), DocxLinksTool())
        ]

    assert [payload["tool"] for payload in payloads] == ["docx_comments", "docx_changes", "docx_notes", "docx_media", "docx_links"]
    assert all(tool.parameters["properties"]["limit"]["maximum"] == 200 for tool in (DocxCommentsTool(), DocxChangesTool(), DocxNotesTool(), DocxMediaTool(), DocxLinksTool()))
    assert {"docx_comments", "docx_changes", "docx_notes", "docx_media", "docx_links"}.issubset({tool.name for tool in builtin_tool_registry.get_all_tools()})


def test_docx_read_schema_requires_explicit_review_scopes_only():
    scopes = DocxReadTool().parameters["properties"]["scopes"]
    assert scopes == {"type": "array", "items": {"type": "string", "enum": ["comments", "changes", "notes"]}, "uniqueItems": True}
