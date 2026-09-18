import logging
import hashlib
import shutil
import subprocess
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement

from app.services import doc_retrieval
from app.services.docx_service import docx_service


def _write_sample_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("A plain paragraph.")
    document.add_paragraph("First checklist item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "苹果"
    table.cell(1, 1).text = "100"
    document.add_paragraph("Closing note.")
    document.save(path)


def _write_merged_table_docx(path: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Merged heading"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Product"
    table.cell(1, 1).text = "Amount"
    document.save(path)


def _write_malformed_xml_docx(path: Path) -> None:
    _write_sample_docx(path)
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["word/document.xml"] = b"<w:document>"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


def _add_package_metadata(path: Path) -> None:
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["docProps/app.xml"] = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">
      <Template>Normal.dotm</Template><Application>Microsoft Office Word</Application>
      <AppVersion>16.0000</AppVersion>
    </Properties>'''
    members["docProps/custom.xml"] = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
      xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
      <property fmtid="{00000000-0000-0000-0000-000000000000}" pid="2" name="Classification"><vt:lpwstr>Internal</vt:lpwstr></property>
      <property fmtid="{00000000-0000-0000-0000-000000000000}" pid="3" name="Retained"><vt:bool>true</vt:bool></property>
    </Properties>'''
    members["word/settings.xml"] = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:documentProtection w:edit="readOnly" w:enforcement="1" w:formatting="0"/>
    </w:settings>'''
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


def test_extract_blocks_preserves_document_order_stable_locators_and_audit(tmp_path: Path, caplog):
    caplog.set_level(logging.INFO)
    path = tmp_path / "sample.docx"
    _write_sample_docx(path)

    result = docx_service.extract_blocks(
        "sample.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["ok"] is True
    assert [block["block_id"] for block in result["blocks"]] == [
        "p-0001",
        "p-0002",
        "p-0003",
        "table-0004",
        "p-0005",
    ]
    assert [block["kind"] for block in result["blocks"]] == [
        "heading",
        "paragraph",
        "list_item",
        "table",
        "paragraph",
    ]
    assert result["blocks"][0]["level"] == 1
    assert result["blocks"][1]["locator"] == {
        "block_id": "p-0002",
        "kind": "paragraph",
        "source_index": 2,
        "heading_path": ["Overview"],
    }
    assert result["blocks"][2]["text"] == "First checklist item"
    assert result["blocks"][3]["rows"][1][0] == {
        "text": "苹果",
        "locator": {
            "block_id": "table-0004",
            "kind": "table_cell",
            "source_index": 4,
            "heading_path": ["Overview"],
            "row": 2,
            "column": 1,
        },
    }
    assert result["fidelity_warnings"] == []
    assert any("[DocxAudit]" in message for message in caplog.messages)


def test_extract_blocks_reports_merged_table_ranges(tmp_path: Path):
    path = tmp_path / "merged.docx"
    _write_merged_table_docx(path)

    result = docx_service.extract_blocks(
        "merged.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["ok"] is True
    assert result["blocks"][0]["merged_cells"] == [
        {
            "start": {"row": 1, "column": 1},
            "end": {"row": 1, "column": 2},
        }
    ]


def test_extract_blocks_reports_unsupported_body_content_as_a_fidelity_warning(tmp_path: Path):
    path = tmp_path / "unsupported.docx"
    document = Document()
    document.add_paragraph("Readable")
    document.element.body.append(OxmlElement("w:altChunk"))
    document.save(path)

    result = docx_service.extract_blocks(
        "unsupported.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["ok"] is True
    assert result["fidelity_warnings"] == [
        {"code": "UNSUPPORTED_BODY_ELEMENT", "element": "altChunk"}
    ]


def test_extract_blocks_rejects_paths_outside_allowed_roots(tmp_path: Path):
    with pytest.raises(doc_retrieval.DocAccessError):
        docx_service.extract_blocks(
            "/etc/passwd",
            root_dir=str(tmp_path),
            allow_roots=[str(tmp_path)],
        )


def test_extract_blocks_enforces_docx_extension_and_deny_roots(tmp_path: Path, caplog):
    caplog.set_level(logging.INFO)
    (tmp_path / "not-docx.txt").write_text("plain text", encoding="utf-8")

    with pytest.raises(doc_retrieval.DocAccessError, match="File extension is not allowed"):
        docx_service.extract_blocks(
            "not-docx.txt",
            root_dir=str(tmp_path),
            allow_roots=[str(tmp_path)],
        )

    sample_path = tmp_path / "sample.docx"
    _write_sample_docx(sample_path)
    with pytest.raises(doc_retrieval.DocAccessError, match="denied"):
        docx_service.extract_blocks(
            "sample.docx",
            root_dir=str(tmp_path),
            allow_roots=[str(tmp_path)],
            deny_roots=[str(tmp_path)],
        )
    assert any("DOCX_ACCESS_DENIED" in message for message in caplog.messages)


def test_extract_blocks_rejects_ambiguous_relative_paths(tmp_path: Path):
    first_root = tmp_path / "first"
    second_root = tmp_path / "second"
    first_root.mkdir()
    second_root.mkdir()
    _write_sample_docx(first_root / "same.docx")
    _write_sample_docx(second_root / "same.docx")

    with pytest.raises(doc_retrieval.DocAccessError, match="Ambiguous relative path"):
        docx_service.extract_blocks(
            "same.docx",
            allow_roots=[str(first_root), str(second_root)],
        )


def test_extract_blocks_reports_malformed_docx_without_parser_traceback(tmp_path: Path):
    (tmp_path / "broken.docx").write_bytes(b"not a DOCX package")

    result = docx_service.extract_blocks(
        "broken.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result == {
        "ok": False,
        "error_code": "DOCX_PARSE_FAILED",
        "message": "The file is not a readable DOCX package.",
        "fidelity_warnings": [],
    }


def test_extract_blocks_normalizes_malformed_ooxml_xml(tmp_path: Path):
    path = tmp_path / "broken-xml.docx"
    _write_malformed_xml_docx(path)

    result = docx_service.extract_blocks(
        "broken-xml.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result == {
        "ok": False,
        "error_code": "DOCX_PARSE_FAILED",
        "message": "The file is not a readable DOCX package.",
        "fidelity_warnings": [],
    }


def test_search_returns_deterministic_block_and_table_cell_citations(tmp_path: Path):
    path = tmp_path / "sample.docx"
    _write_sample_docx(path)

    result = docx_service.search(
        "sample.docx",
        query="product",
        domains=["headings", "table_cells"],
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["ok"] is True
    assert result["total_matches"] == 1
    assert result["matches"] == [
        {
            "text": "Product",
            "snippet": "Product",
            "match_spans": [{"start": 0, "end": 7}],
            "score": 1.0,
            "citation": {
                "block_id": "table-0004",
                "kind": "table_cell",
                "source_index": 4,
                "heading_path": ["Overview"],
                "row": 1,
                "column": 1,
            },
        }
    ]
    assert result["is_truncated"] is False


def test_query_applies_constrained_filters_without_executing_expression(tmp_path: Path):
    path = tmp_path / "sample.docx"
    _write_sample_docx(path)

    result = docx_service.query(
        "sample.docx",
        query={"kind": "paragraph", "heading_path": ["Overview"], "text": "plain"},
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["ok"] is True
    assert result["applied_filters"] == {
        "kind": "paragraph",
        "heading_path": ["Overview"],
        "text": "plain",
    }
    assert [item["block_id"] for item in result["results"]] == ["p-0002"]
    invalid = docx_service.query(
        "sample.docx",
        query={"expression": "__import__('os').system('false')"},
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )
    assert invalid == {
        "ok": False,
        "error_code": "DOCX_QUERY_INVALID",
        "message": "Unsupported query field: expression",
    }


def test_structure_exposes_headings_lists_tables_and_section_boundaries(tmp_path: Path):
    path = tmp_path / "sample.docx"
    _write_sample_docx(path)

    result = docx_service.structure(
        "sample.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["ok"] is True
    assert result["headings"] == [
        {
            "block_id": "p-0001",
            "text": "Overview",
            "level": 1,
            "heading_path": ["Overview"],
            "block_range": {"start": 1, "end": 5},
        }
    ]
    assert result["lists"] == [{"block_id": "p-0003", "level": 1, "text": "First checklist item"}]
    assert result["tables"] == [{"table_id": "table-0004", "source_index": 4, "heading_path": ["Overview"]}]
    assert result["sections"] == [{"index": 1, "block_range": {"start": 1, "end": 5}}]
    assert result["headers_footers"] == {"has_headers": False, "has_footers": False}


def test_structure_reports_distinct_section_block_ranges(tmp_path: Path):
    path = tmp_path / "sections.docx"
    document = Document()
    document.add_paragraph("First section")
    document.add_section()
    document.add_paragraph("Second section")
    document.save(path)

    result = docx_service.structure(
        "sections.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["sections"] == [
        {"index": 1, "block_range": {"start": 1, "end": 2}},
        {"index": 2, "block_range": {"start": 3, "end": 3}},
    ]
    query = docx_service.query(
        "sections.docx",
        query={"section": 2},
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )
    assert [item["block_id"] for item in query["results"]] == ["p-0003"]


def test_structure_reports_ooxml_list_nesting_level(tmp_path: Path):
    path = tmp_path / "nested-list.docx"
    document = Document()
    nested_item = document.add_paragraph("Nested item", style="List Bullet")
    number_properties = OxmlElement("w:numPr")
    nesting_level = OxmlElement("w:ilvl")
    nesting_level.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "1")
    number_properties.append(nesting_level)
    nested_item._p.get_or_add_pPr().append(number_properties)
    document.save(path)

    result = docx_service.structure(
        "nested-list.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["lists"] == [{"block_id": "p-0001", "level": 2, "text": "Nested item"}]


def test_structure_anchors_bookmarks_to_their_containing_block(tmp_path: Path):
    path = tmp_path / "bookmark.docx"
    document = Document()
    paragraph = document.add_paragraph("Anchored content")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id", "7")
    bookmark.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name", "target")
    paragraph._p.append(bookmark)
    document.save(path)

    result = docx_service.structure(
        "bookmark.docx",
        root_dir=str(tmp_path),
        allow_roots=[str(tmp_path)],
    )

    assert result["bookmarks"] == [{
        "name": "target", "id": "7",
        "citation": {"block_id": "p-0001", "kind": "paragraph", "source_index": 1, "heading_path": []},
    }]


def test_metadata_reads_package_properties_without_reading_body(tmp_path: Path, monkeypatch, caplog):
    caplog.set_level(logging.INFO)
    path = tmp_path / "metadata.docx"
    document = Document()
    document.core_properties.title = "Metadata brief"
    document.core_properties.author = "Yue"
    document.core_properties.subject = "Release"
    document.core_properties.revision = 4
    document.core_properties.language = "en-US"
    document.save(path)
    _add_package_metadata(path)
    source_digest = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr("app.services.docx_service.Document", lambda *_: pytest.fail("metadata must not load body"))

    result = docx_service.metadata("metadata.docx", root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result["ok"] is True
    assert result["core"]["title"] == "Metadata brief"
    assert result["core"]["author"] == "Yue"
    assert result["core"]["subject"] == "Release"
    assert result["core"]["revision"] == 4
    assert result["core"]["language"] == "en-US"
    assert result["application"] == {
        "application": "Microsoft Office Word",
        "template": "Normal.dotm",
        "app_version": "16.0000",
    }
    assert result["custom"] == {"Classification": "Internal", "Retained": True}
    assert result["protection"] == {"enabled": True, "edit": "readOnly", "enforcement": True, "formatting": False}
    assert hashlib.sha256(path.read_bytes()).hexdigest() == source_digest
    assert any("[DocxAudit]" in message and "docx_metadata" in message for message in caplog.messages)


def test_render_returns_actionable_error_without_a_renderer(tmp_path: Path, monkeypatch):
    path = tmp_path / "render.docx"
    Document().save(path)
    monkeypatch.setattr(docx_service, "_renderer_path", lambda: None, raising=False)

    result = docx_service.render("render.docx", root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result == {"ok": False, "error_code": "DOCX_RENDERER_UNAVAILABLE", "message": "DOCX rendering requires LibreOffice (soffice)."}


def test_render_rejects_invalid_range_before_checking_renderer(tmp_path: Path, monkeypatch):
    path = tmp_path / "render.docx"
    Document().save(path)
    monkeypatch.setattr(docx_service, "_renderer_path", lambda: None)

    result = docx_service.render("render.docx", page_start=3, page_end=2, root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result == {"ok": False, "error_code": "DOCX_RENDER_INVALID", "message": "page_end must be greater than or equal to page_start."}


def test_render_returns_page_artifacts_with_renderer_provenance(tmp_path: Path, monkeypatch):
    path = tmp_path / "render.docx"
    Document().save(path)
    source_digest = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr(docx_service, "_renderer_path", lambda: "/mock/soffice")
    monkeypatch.setattr(
        docx_service,
        "_renderer_version",
        lambda command: "LibreOffice 25.2" if command == "/mock/soffice" else "unavailable",
    )
    monkeypatch.setattr("app.services.docx_service.shutil.which", lambda name: "/mock/pdftoppm" if name == "pdftoppm" else None)

    def fake_run(command, **kwargs):
        if command[0] == "/mock/soffice":
            output_dir = Path(command[command.index("--outdir") + 1])
            (output_dir / "render.pdf").write_bytes(b"%PDF-1.4")
        else:
            prefix = Path(command[-1])
            for page in range(int(command[command.index("-f") + 1]), int(command[command.index("-l") + 1]) + 1):
                prefix.parent.joinpath(f"{prefix.name}-{page}.png").write_bytes(b"PNG")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr("app.services.docx_service.subprocess.run", fake_run)

    result = docx_service.render("render.docx", page_start=2, page_end=3, root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result["ok"] is True
    assert [artifact["page"] for artifact in result["artifacts"]] == [2, 3]
    assert all(Path(artifact["path"]).is_file() for artifact in result["artifacts"])
    assert result["provenance"] == {
        "backend": {"name": "LibreOffice", "version": "LibreOffice 25.2"},
        "image_converter": {"name": "pdftoppm", "version": "unavailable"},
        "page_to_source": {"status": "unavailable"},
    }
    assert hashlib.sha256(path.read_bytes()).hexdigest() == source_digest


def test_render_never_returns_pages_when_conversion_does_not_create_them(tmp_path: Path, monkeypatch):
    path = tmp_path / "render.docx"
    Document().save(path)
    monkeypatch.setattr(docx_service, "_renderer_path", lambda: "/mock/soffice")
    monkeypatch.setattr("app.services.docx_service.shutil.which", lambda name: "/mock/pdftoppm" if name == "pdftoppm" else None)

    def fake_run(command, **kwargs):
        if "--convert-to" in command:
            Path(command[command.index("--outdir") + 1], "render.pdf").write_bytes(b"%PDF-1.4")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr("app.services.docx_service.subprocess.run", fake_run)

    result = docx_service.render("render.docx", root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result == {"ok": False, "error_code": "DOCX_RENDER_FAILED", "message": "Image conversion did not produce the requested pages."}


def test_render_refuses_page_results_when_renderer_versions_are_unavailable(tmp_path: Path, monkeypatch):
    path = tmp_path / "render.docx"
    Document().save(path)
    monkeypatch.setattr(docx_service, "_renderer_path", lambda: "/mock/soffice")
    monkeypatch.setattr("app.services.docx_service.shutil.which", lambda name: "/mock/pdftoppm" if name == "pdftoppm" else None)
    monkeypatch.setattr(docx_service, "_renderer_version", lambda _: None)

    def fake_run(command, **kwargs):
        if "--convert-to" in command:
            Path(command[command.index("--outdir") + 1], "render.pdf").write_bytes(b"%PDF-1.4")
        elif "-png" in command:
            Path(f"{command[-1]}-1.png").write_bytes(b"PNG")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr("app.services.docx_service.subprocess.run", fake_run)

    result = docx_service.render("render.docx", root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result == {
        "ok": False,
        "error_code": "DOCX_RENDER_PROVENANCE_UNAVAILABLE",
        "message": "Renderer version provenance is unavailable; no page artifacts were returned.",
    }


def test_render_smoke_produces_inspection_artifacts_without_modifying_source(tmp_path: Path):
    if not docx_service._renderer_path() or not shutil.which("pdftoppm"):
        pytest.skip("LibreOffice and pdftoppm are required for the render smoke test")
    path = tmp_path / "render-smoke.docx"
    document = Document()
    document.add_paragraph("First page")
    document.add_page_break()
    document.add_paragraph("Second page")
    document.save(path)
    source_digest = hashlib.sha256(path.read_bytes()).hexdigest()

    result = docx_service.render("render-smoke.docx", page_start=1, page_end=2, root_dir=str(tmp_path), allow_roots=[str(tmp_path)])

    assert result["ok"] is True
    assert [artifact["page"] for artifact in result["artifacts"]] == [1, 2]
    assert all(Path(artifact["path"]).is_file() for artifact in result["artifacts"])
    assert result["provenance"]["backend"]["version"] != "unavailable"
    assert all(artifact["provenance"] == result["provenance"] for artifact in result["artifacts"])
    assert hashlib.sha256(path.read_bytes()).hexdigest() == source_digest
