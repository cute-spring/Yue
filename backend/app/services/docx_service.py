"""Read-only DOCX package primitives shared by future MCP tools."""

import logging
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Any, Dict, List, Optional
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from lxml.etree import XMLSyntaxError

from . import doc_retrieval


DOCX_EXTENSIONS = [".docx"]
READ_BLOCK_LIMIT = 500
TABLE_LIMIT = 100
TABLE_ROW_CELL_LIMIT = 500
RETRIEVAL_RESULT_LIMIT = 200
RENDER_PAGE_LIMIT = 20
REVIEW_RESULT_LIMIT = 200
SEARCH_DOMAINS = {"body", "headings", "table_cells"}
QUERY_FIELDS = {"kind", "heading_path", "style", "section", "table_id", "text", "limit"}
READ_SCOPES = {"comments", "changes", "notes"}
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"w": WORD_NS, "r": REL_NS}
logger = logging.getLogger(__name__)


class DocxService:
    """Expose a small, safe interface for reading ordered DOCX body blocks."""

    def _resolve_path(
        self,
        path: str,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> str:
        docs_root = doc_retrieval.resolve_docs_root_for_read(
            path,
            requested_root=root_dir,
            allow_roots=allow_roots,
            deny_roots=deny_roots,
            allowed_extensions=DOCX_EXTENSIONS,
            require_md=False,
        )
        return doc_retrieval.resolve_docs_path(
            path,
            docs_root=docs_root,
            require_md=False,
            allowed_extensions=DOCX_EXTENSIONS,
        )

    def extract_blocks(
        self,
        path: str,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Return body content in document order with stable, local locators."""
        started_at = time.monotonic()
        try:
            abs_path = self._resolve_path(path, root_dir, allow_roots, deny_roots)
        except doc_retrieval.DocAccessError:
            self._log_audit(path, started_at, status="failed", error_code="DOCX_ACCESS_DENIED")
            raise
        try:
            document = Document(abs_path)
        except (BadZipFile, KeyError, OSError, PackageNotFoundError, ValueError, XMLSyntaxError):
            self._log_audit(path, started_at, status="failed", error_code="DOCX_PARSE_FAILED")
            return {
                "ok": False,
                "error_code": "DOCX_PARSE_FAILED",
                "message": "The file is not a readable DOCX package.",
                "fidelity_warnings": [],
            }

        blocks: List[Dict[str, Any]] = []
        fidelity_warnings: List[Dict[str, str]] = []
        heading_path: List[str] = []
        for index, element in enumerate(document.element.body.iterchildren(), start=1):
            if isinstance(element, CT_P):
                block = self._paragraph_block(Paragraph(element, document), index)
                if block["kind"] == "heading":
                    heading_path = heading_path[: block["level"] - 1]
                    heading_path.append(block["text"])
                block["locator"] = self._locator(
                    block["block_id"], block["kind"], index, heading_path
                )
                blocks.append(block)
            elif isinstance(element, CT_Tbl):
                blocks.append(self._table_block(Table(element, document), index, heading_path))
            elif element.tag.rsplit("}", maxsplit=1)[-1] != "sectPr":
                fidelity_warnings.append(
                    {
                        "code": "UNSUPPORTED_BODY_ELEMENT",
                        "element": element.tag.rsplit("}", maxsplit=1)[-1],
                    }
                )

        result = {
            "ok": True,
            "file": path,
            "blocks": blocks,
            "fidelity_warnings": fidelity_warnings,
            "elapsed_ms": round((time.monotonic() - started_at) * 1000, 2),
        }
        self._log_audit(
            path,
            started_at,
            blocks=len(blocks),
            fidelity_warnings=len(fidelity_warnings),
        )
        return result

    def profile(
        self,
        path: str,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        blocks_result = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not blocks_result["ok"]:
            return blocks_result
        document = Document(self._resolve_path(path, root_dir, allow_roots, deny_roots))
        blocks = blocks_result["blocks"]
        return {
            "ok": True,
            "file": path,
            "metadata": {
                "title": document.core_properties.title or "",
                "author": document.core_properties.author or "",
            },
            "counts": {
                "headings": sum(block["kind"] == "heading" for block in blocks),
                "paragraphs": sum(block["kind"] in {"heading", "paragraph", "list_item"} for block in blocks),
                "tables": sum(block["kind"] == "table" for block in blocks),
                "sections": len(document.sections),
            },
            "heading_outline": [{"text": block["text"], "level": block["level"], "block_id": block["block_id"]} for block in blocks if block["kind"] == "heading"],
            "features": {"has_tables": any(block["kind"] == "table" for block in blocks), "has_fidelity_warnings": bool(blocks_result["fidelity_warnings"])},
            "fidelity_warnings": blocks_result["fidelity_warnings"],
        }

    def read(
        self,
        path: str,
        cursor: int = 0,
        limit: int = 200,
        mode: str = "json",
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
        scopes: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        requested_scopes = set(scopes or [])
        unsupported = sorted(requested_scopes - READ_SCOPES)
        if unsupported:
            return {"ok": False, "error_code": "DOCX_READ_INVALID", "message": f"Unsupported read scope: {unsupported[0]}", "hint": "Use comments, changes, or notes."}
        result = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not result["ok"]:
            return result
        start = max(cursor, 0)
        end = min(start + min(max(limit, 1), READ_BLOCK_LIMIT), len(result["blocks"]))
        blocks = result["blocks"][start:end]
        is_truncated = end < len(result["blocks"])
        response = {
            "ok": True,
            "file": path,
            "data": self._markdown(blocks) if mode == "markdown" else blocks,
            "is_truncated": is_truncated,
            "next_cursor": end if is_truncated else None,
            "citations": [block["locator"] for block in blocks],
            "fidelity_warnings": result["fidelity_warnings"],
        }
        if requested_scopes:
            response["layers"] = {
                scope: getattr(self, scope)(path, limit=limit, root_dir=root_dir, allow_roots=allow_roots, deny_roots=deny_roots)
                for scope in sorted(requested_scopes)
            }
        return response

    def extract_tables(
        self,
        path: str,
        table_id: Optional[str] = None,
        table_index: Optional[int] = None,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        result = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not result["ok"]:
            return result
        tables = [block for block in result["blocks"] if block["kind"] == "table"]
        if table_id:
            tables = [table for table in tables if table["table_id"] == table_id]
        elif table_index is not None:
            tables = tables[table_index - 1:table_index]
        is_truncated = len(tables) > TABLE_LIMIT
        tables = tables[:TABLE_LIMIT]
        bounded_tables = []
        for table in tables:
            bounded_rows = []
            for row in table["rows"]:
                if len(row) > TABLE_ROW_CELL_LIMIT:
                    is_truncated = True
                bounded_rows.append(row[:TABLE_ROW_CELL_LIMIT])
            bounded_tables.append({**table, "rows": bounded_rows})
        return {
            "ok": True,
            "file": path,
            "tables": bounded_tables,
            "is_truncated": is_truncated,
            "fidelity_warnings": result["fidelity_warnings"],
        }

    def search(
        self,
        path: str,
        query: str,
        domains: Optional[List[str]] = None,
        regex: bool = False,
        limit: int = 50,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Search a fixed, visible DOCX content domain with stable citations."""
        result = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not result["ok"]:
            return result
        requested_domains = set(domains or ["body", "headings", "table_cells"])
        unsupported = sorted(requested_domains - SEARCH_DOMAINS)
        if unsupported:
            return {
                "ok": False,
                "error_code": "DOCX_SEARCH_INVALID",
                "message": f"Unsupported search domain: {unsupported[0]}",
            }
        try:
            matcher = re.compile(query if regex else re.escape(query), flags=re.IGNORECASE)
        except re.error as error:
            return {"ok": False, "error_code": "DOCX_SEARCH_INVALID", "message": f"Invalid regex: {error}"}
        if not query:
            return {"ok": False, "error_code": "DOCX_SEARCH_INVALID", "message": "Search query must not be empty."}

        matches: List[Dict[str, Any]] = []
        for text, citation, domain in self._searchable_items(result["blocks"]):
            if domain not in requested_domains:
                continue
            spans = [{"start": match.start(), "end": match.end()} for match in matcher.finditer(text)]
            if not spans:
                continue
            matches.append(
                {
                    "text": text,
                    "snippet": self._snippet(text, spans[0]),
                    "match_spans": spans,
                    "score": 1.0,
                    "citation": citation,
                }
            )
        bounded_limit = min(max(limit, 1), RETRIEVAL_RESULT_LIMIT)
        return {
            "ok": True,
            "file": path,
            "query": query,
            "domains": sorted(requested_domains),
            "matches": matches[:bounded_limit],
            "total_matches": len(matches),
            "is_truncated": len(matches) > bounded_limit,
            "limit": bounded_limit,
            "fidelity_warnings": result["fidelity_warnings"],
        }

    def comments(self, path: str, limit: int = 50, author: Optional[str] = None, root_dir: Optional[str] = None, allow_roots: Optional[List[str]] = None, deny_roots: Optional[List[str]] = None) -> Dict[str, Any]:
        """Read comment metadata and anchors without changing the package."""
        context = self._review_context(path, root_dir, allow_roots, deny_roots)
        if not context["ok"]:
            return context
        root = context["parts"].get("word/comments.xml")
        extensions = context["parts"].get("word/commentsExtended.xml")
        extension_by_para = {item.get("{http://schemas.microsoft.com/office/word/2012/wordml}paraId", ""): item for item in (list(extensions) if extensions is not None else [])}
        comments = []
        if root is not None:
            for item in root.findall("w:comment", NSMAP):
                comment: Dict[str, Any] = {"id": item.get(f"{{{WORD_NS}}}id", "")}
                for attribute, key in (("author", "author"), ("date", "date")):
                    value = item.get(f"{{{WORD_NS}}}{attribute}")
                    if value:
                        comment[key] = value
                comment["text"] = self._xml_text(item)
                para_id = item.get("{http://schemas.microsoft.com/office/word/2012/wordml}paraId", "")
                extension = extension_by_para.get(para_id)
                if extension is not None:
                    parent = extension.get("{http://schemas.microsoft.com/office/word/2012/wordml}paraIdParent")
                    if parent:
                        comment["reply_to"] = parent
                    done = extension.get("{http://schemas.microsoft.com/office/word/2012/wordml}done")
                    if done is not None:
                        comment["status"] = "resolved" if done in {"1", "true", "True"} else "open"
                citation = context["comment_anchors"].get(comment["id"])
                if citation:
                    comment["citation"] = citation
                if author is None or comment.get("author") == author:
                    comments.append(comment)
        return self._bounded_review_result(path, "comments", comments, limit, context["started_at"])

    def changes(self, path: str, limit: int = 50, view: str = "markup", author: Optional[str] = None, root_dir: Optional[str] = None, allow_roots: Optional[List[str]] = None, deny_roots: Optional[List[str]] = None) -> Dict[str, Any]:
        """Return tracked insertions and deletions as inert markup records."""
        context = self._review_context(path, root_dir, allow_roots, deny_roots)
        if not context["ok"]:
            return context
        if view not in {"final", "original", "markup"}:
            return {"ok": False, "error_code": "DOCX_CHANGES_INVALID", "message": "Unsupported change view.", "hint": "Use final, original, or markup."}
        changes = []
        for element in context["document"].xpath(".//w:ins | .//w:del | .//w:rPrChange | .//w:pPrChange", namespaces=NSMAP):
            local_name = etree.QName(element).localname
            change: Dict[str, Any] = {
                "id": element.get(f"{{{WORD_NS}}}id", ""),
                "kind": {"ins": "insertion", "del": "deletion"}.get(local_name, "formatting"),
            }
            for attribute, key in (("author", "author"), ("date", "date")):
                value = element.get(f"{{{WORD_NS}}}{attribute}")
                if value:
                    change[key] = value
            change["text"] = self._xml_text(element)
            citation = self._node_citation(element, context["body_children"], context["locators"])
            if citation:
                change["citation"] = citation
            if author is None or change.get("author") == author:
                changes.append(change)
        if view == "final":
            changes = [change for change in changes if change["kind"] != "deletion"]
        elif view == "original":
            changes = [change for change in changes if change["kind"] != "insertion"]
        return {**self._bounded_review_result(path, "changes", changes, limit, context["started_at"]), "view": view}

    def notes(self, path: str, limit: int = 50, kind: Optional[str] = None, root_dir: Optional[str] = None, allow_roots: Optional[List[str]] = None, deny_roots: Optional[List[str]] = None) -> Dict[str, Any]:
        """Return footnotes and endnotes with their body references when present."""
        context = self._review_context(path, root_dir, allow_roots, deny_roots)
        if not context["ok"]:
            return context
        if kind not in {None, "footnote", "endnote"}:
            return {"ok": False, "error_code": "DOCX_NOTES_INVALID", "message": "Unsupported note kind.", "hint": "Use footnote or endnote."}
        notes = []
        for kind, member, tag in (("footnote", "word/footnotes.xml", "footnote"), ("endnote", "word/endnotes.xml", "endnote")):
            root = context["parts"].get(member)
            if root is None:
                continue
            for element in root.findall(f"w:{tag}", NSMAP):
                note = {"id": element.get(f"{{{WORD_NS}}}id", ""), "kind": kind, "text": self._xml_text(element)}
                citation = context["note_anchors"].get((kind, note["id"]))
                if citation:
                    note["citation"] = citation
                if kind is None or note["kind"] == kind:
                    notes.append(note)
        return self._bounded_review_result(path, "notes", notes, limit, context["started_at"])

    def media(self, path: str, limit: int = 50, root_dir: Optional[str] = None, allow_roots: Optional[List[str]] = None, deny_roots: Optional[List[str]] = None) -> Dict[str, Any]:
        """Inventory package media by name only; payloads are never extracted or opened."""
        context = self._review_context(path, root_dir, allow_roots, deny_roots)
        if not context["ok"]:
            return context
        items = []
        for member in context["members"]:
            if member.startswith("word/media/"):
                item: Dict[str, Any] = {"kind": "image", "path": member, "mime_type": self._media_mime(member), "extracted": False}
                anchor = self._relationship_anchor(context, member)
                if anchor:
                    item.update(anchor)
                items.append(item)
            elif member.startswith("word/charts/"):
                items.append({"kind": "chart", "path": member, "mime_type": "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", **(self._relationship_anchor(context, member) or {}), "extracted": False})
            elif member.startswith("word/embeddings/"):
                items.append({"kind": "embedded_file", "path": member, "mime_type": "application/octet-stream", **(self._relationship_anchor(context, member) or {}), "extracted": False})
        items.sort(key=lambda item: {"image": 0, "chart": 1, "embedded_file": 2}[item["kind"]])
        return self._bounded_review_result(path, "items", items, limit, context["started_at"])

    def links(self, path: str, limit: int = 50, root_dir: Optional[str] = None, allow_roots: Optional[List[str]] = None, deny_roots: Optional[List[str]] = None) -> Dict[str, Any]:
        """Inventory links and references without following any external target."""
        context = self._review_context(path, root_dir, allow_roots, deny_roots)
        if not context["ok"]:
            return context
        links = []
        for hyperlink in context["document"].xpath(".//w:hyperlink", namespaces=NSMAP):
            citation = self._node_citation(hyperlink, context["body_children"], context["locators"])
            rel_id = hyperlink.get(f"{{{REL_NS}}}id")
            if rel_id and rel_id in context["relationships"]:
                target = context["relationships"][rel_id].get("target")
                links.append({"kind": "external", "destination": target, "followed": False, **({"citation": citation} if citation else {})})
            elif hyperlink.get(f"{{{WORD_NS}}}anchor"):
                links.append({"kind": "internal_anchor", "anchor": hyperlink.get(f"{{{WORD_NS}}}anchor"), **({"citation": citation} if citation else {})})
        for bookmark in context["document"].xpath(".//w:bookmarkStart", namespaces=NSMAP):
            citation = self._node_citation(bookmark, context["body_children"], context["locators"])
            links.append({"kind": "bookmark", "anchor": bookmark.get(f"{{{WORD_NS}}}name", ""), **({"citation": citation} if citation else {})})
        for field in context["document"].xpath(".//w:fldSimple", namespaces=NSMAP):
            match = re.search(r"\bREF\s+([^\s]+)", field.get(f"{{{WORD_NS}}}instr", ""))
            if match:
                citation = self._node_citation(field, context["body_children"], context["locators"])
                links.append({"kind": "cross_reference", "anchor": match.group(1), **({"citation": citation} if citation else {})})
        result = self._bounded_review_result(path, "links", links, limit, context["started_at"])
        relationship_items = [
            {"id": identifier, **relationship, "followed": False}
            for identifier, relationship in sorted(context["relationships"].items())
        ]
        result["relationships"] = relationship_items[:REVIEW_RESULT_LIMIT]
        if len(relationship_items) > REVIEW_RESULT_LIMIT:
            result["is_truncated"] = True
            result["next_cursor"] = "narrow relationship inspection by link type"
        return result

    def _review_context(self, path: str, root_dir: Optional[str], allow_roots: Optional[List[str]], deny_roots: Optional[List[str]]) -> Dict[str, Any]:
        started_at = time.monotonic()
        blocks = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not blocks["ok"]:
            return blocks
        try:
            with ZipFile(self._resolve_path(path, root_dir, allow_roots, deny_roots)) as package:
                members = package.namelist()
                parts = {member: self._package_xml(package, member) for member in ("word/document.xml", "word/comments.xml", "word/commentsExtended.xml", "word/footnotes.xml", "word/endnotes.xml", "word/_rels/document.xml.rels")}
        except (BadZipFile, OSError, ValueError, XMLSyntaxError, etree.LxmlError):
            self._log_audit(path, started_at, status="failed", tool="docx_review", error_code="DOCX_PARSE_FAILED")
            return {"ok": False, "error_code": "DOCX_PARSE_FAILED", "message": "The file is not a readable DOCX package.", "hint": "Verify the file is an unencrypted DOCX package."}
        document = parts["word/document.xml"]
        if document is None:
            self._log_audit(path, started_at, status="failed", tool="docx_review", error_code="DOCX_PARSE_FAILED")
            return {"ok": False, "error_code": "DOCX_PARSE_FAILED", "message": "The file is not a readable DOCX package.", "hint": "Verify the file is an unencrypted DOCX package."}
        body = document.find(f"{{{WORD_NS}}}body")
        body_children = list(body) if body is not None else []
        locators = {block["locator"]["source_index"]: block["locator"] for block in blocks["blocks"]}
        relationship_root = parts["word/_rels/document.xml.rels"]
        relationships = {
            item.get("Id", ""): {"target": item.get("Target", ""), "target_mode": item.get("TargetMode", ""), "type": item.get("Type", "")}
            for item in (list(relationship_root) if relationship_root is not None else [])
        }
        comment_anchors, note_anchors = {}, {}
        for element in document.xpath(".//w:commentRangeStart | .//w:footnoteReference | .//w:endnoteReference", namespaces=NSMAP):
            citation = self._node_citation(element, body_children, locators)
            if citation is None:
                continue
            name, identifier = etree.QName(element).localname, element.get(f"{{{WORD_NS}}}id", "")
            if name == "commentRangeStart":
                comment_anchors[identifier] = citation
            else:
                note_anchors[("footnote" if name == "footnoteReference" else "endnote", identifier)] = citation
        return {"ok": True, "started_at": started_at, "parts": parts, "members": members, "document": document, "body_children": body_children, "locators": locators, "relationships": relationships, "comment_anchors": comment_anchors, "note_anchors": note_anchors}

    def _bounded_review_result(self, path: str, key: str, items: List[Dict[str, Any]], limit: int, started_at: float) -> Dict[str, Any]:
        bounded_limit = min(max(int(limit), 1), REVIEW_RESULT_LIMIT)
        result = {"ok": True, "file": path, key: items[:bounded_limit], "is_truncated": len(items) > bounded_limit, "limit": bounded_limit}
        self._log_audit(path, started_at, tool={"items": "docx_media"}.get(key, f"docx_{key.rstrip('s')}"), count=len(result[key]), is_truncated=result["is_truncated"])
        return result

    @staticmethod
    def _xml_text(element: Any) -> str:
        return "".join(element.xpath(".//w:t/text() | .//w:delText/text()", namespaces=NSMAP))

    @staticmethod
    def _node_citation(element: Any, body_children: List[Any], locators: Dict[int, Dict[str, Any]]) -> Optional[Dict[str, Any]]:
        owner = next((ancestor for ancestor in (element, *element.iterancestors()) if ancestor in body_children), None)
        return locators.get(body_children.index(owner) + 1) if owner is not None else None

    @staticmethod
    def _media_mime(member: str) -> str:
        return {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".svg": "image/svg+xml"}.get(Path(member).suffix.lower(), "application/octet-stream")

    def _relationship_anchor(self, context: Dict[str, Any], member: str) -> Optional[Dict[str, Any]]:
        rel_id = next((identifier for identifier, relationship in context["relationships"].items() if relationship["target"] == member.removeprefix("word/")), None)
        if rel_id is None:
            return None
        node = next((element for element in context["document"].iter() if rel_id in element.attrib.values()), None)
        if node is None:
            return None
        citation = self._node_citation(node, context["body_children"], context["locators"])
        result: Dict[str, Any] = {}
        drawing_ns = {**NSMAP, "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
        doc_properties = next(iter(node.xpath("ancestor::w:drawing//wp:docPr | .//wp:docPr", namespaces=drawing_ns)), None)
        extent = next(iter(node.xpath("ancestor::w:drawing//wp:extent | .//wp:extent", namespaces=drawing_ns)), None)
        if doc_properties is not None and doc_properties.get("descr"):
            result["alt_text"] = doc_properties.get("descr")
        if extent is not None:
            result["dimensions"] = {"cx": int(extent.get("cx", "0")), "cy": int(extent.get("cy", "0"))}
        if citation:
            result["citation"] = citation
        return result

    def metadata(
        self,
        path: str,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Read package properties without extracting document body blocks."""
        started_at = time.monotonic()
        try:
            source = self._resolve_path(path, root_dir, allow_roots, deny_roots)
        except doc_retrieval.DocAccessError:
            self._log_audit(path, started_at, status="failed", tool="docx_metadata", error_code="DOCX_ACCESS_DENIED")
            raise
        try:
            with ZipFile(source) as package:
                core = self._package_core_properties(package)
                application = self._package_application_properties(package)
                custom = self._package_custom_properties(package)
                protection = self._package_protection(package)
        except (BadZipFile, KeyError, OSError, ValueError, XMLSyntaxError, etree.LxmlError):
            self._log_audit(path, started_at, status="failed", tool="docx_metadata", error_code="DOCX_PARSE_FAILED")
            return {"ok": False, "error_code": "DOCX_PARSE_FAILED", "message": "The file is not a readable DOCX package."}
        result = {
            "ok": True,
            "file": path,
            "core": core,
            "application": application,
            "custom": custom,
            "protection": protection,
        }
        self._log_audit(path, started_at, tool="docx_metadata", core_fields=len(core), custom_fields=len(custom))
        return result

    def render(
        self,
        path: str,
        page_start: int = 1,
        page_end: Optional[int] = None,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Render a bounded DOCX page range through LibreOffice and Poppler."""
        started_at = time.monotonic()
        if page_start < 1:
            return self._render_error(path, started_at, "DOCX_RENDER_INVALID", "page_start must be at least 1.")
        if page_end is not None and page_end < page_start:
            return self._render_error(path, started_at, "DOCX_RENDER_INVALID", "page_end must be greater than or equal to page_start.")
        final_page = page_end or page_start
        if final_page - page_start + 1 > RENDER_PAGE_LIMIT:
            return self._render_error(path, started_at, "DOCX_RENDER_INVALID", f"Render requests are limited to {RENDER_PAGE_LIMIT} pages.")
        try:
            source = self._resolve_path(path, root_dir, allow_roots, deny_roots)
        except doc_retrieval.DocAccessError:
            self._log_audit(path, started_at, status="failed", tool="docx_render", error_code="DOCX_ACCESS_DENIED")
            raise
        renderer = self._renderer_path()
        if not renderer:
            return self._render_error(path, started_at, "DOCX_RENDERER_UNAVAILABLE", "DOCX rendering requires LibreOffice (soffice).")
        converter = shutil.which("pdftoppm")
        if not converter:
            return self._render_error(path, started_at, "DOCX_RENDERER_UNAVAILABLE", "DOCX rendering requires pdftoppm for page images.")
        try:
            output_dir = Path(tempfile.mkdtemp(prefix="yue-docx-render-"))
            profile_dir = output_dir / "libreoffice-profile"
            profile_dir.mkdir()
            completed = subprocess.run(
                [
                    renderer,
                    "--headless",
                    "--norestore",
                    "--nodefault",
                    "--nolockcheck",
                    f"-env:UserInstallation={profile_dir.as_uri()}",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output_dir),
                    source,
                ],
                capture_output=True,
                text=True,
                timeout=60,
            )
            pdf = output_dir / f"{Path(source).stem}.pdf"
            if completed.returncode or not pdf.exists():
                return self._render_error(path, started_at, "DOCX_RENDER_FAILED", "LibreOffice could not render this DOCX.")
            prefix = output_dir / "page"
            converted = subprocess.run(
                [converter, "-png", "-f", str(page_start), "-l", str(final_page), str(pdf), str(prefix)],
                capture_output=True,
                text=True,
                timeout=60,
            )
            if converted.returncode:
                return self._render_error(path, started_at, "DOCX_RENDER_FAILED", "Poppler could not convert rendered pages to PNG images.")
            pages = {int(item.stem.rsplit("-", 1)[-1]): item for item in output_dir.glob("page-*.png")}
            requested_pages = list(range(page_start, final_page + 1))
            if any(page not in pages for page in requested_pages):
                return self._render_error(path, started_at, "DOCX_RENDER_FAILED", "Image conversion did not produce the requested pages.")
            renderer_version = self._renderer_version(renderer)
            converter_version = self._renderer_version(converter)
            if not renderer_version or not converter_version:
                return self._render_error(
                    path,
                    started_at,
                    "DOCX_RENDER_PROVENANCE_UNAVAILABLE",
                    "Renderer version provenance is unavailable; no page artifacts were returned.",
                )
            provenance = {
                "backend": {"name": "LibreOffice", "version": renderer_version},
                "image_converter": {"name": "pdftoppm", "version": converter_version},
                "page_to_source": {"status": "unavailable"},
            }
            artifacts = [
                {
                    "page": page,
                    "path": str(pages[page]),
                    "mime_type": "image/png",
                    "temporary": True,
                    "provenance": provenance,
                }
                for page in requested_pages
            ]
            result = {"ok": True, "file": path, "artifacts": artifacts, "is_truncated": False, "provenance": provenance}
            self._log_audit(path, started_at, tool="docx_render", pages=len(artifacts), backend="LibreOffice")
            return result
        except (OSError, subprocess.SubprocessError):
            return self._render_error(path, started_at, "DOCX_RENDER_FAILED", "DOCX rendering could not be completed.")

    def _render_error(self, path: str, started_at: float, error_code: str, message: str) -> Dict[str, Any]:
        self._log_audit(path, started_at, status="failed", tool="docx_render", error_code=error_code)
        return {"ok": False, "error_code": error_code, "message": message}

    @staticmethod
    def _package_xml(package: ZipFile, member: str) -> Optional[Any]:
        try:
            return etree.fromstring(package.read(member))
        except KeyError:
            return None

    @classmethod
    def _package_core_properties(cls, package: ZipFile) -> Dict[str, Any]:
        root = cls._package_xml(package, "docProps/core.xml")
        if root is None:
            return {}
        field_names = {
            "title": "title", "subject": "subject", "creator": "author", "keywords": "keywords",
            "description": "comments", "lastModifiedBy": "last_modified_by", "revision": "revision",
            "category": "category", "contentStatus": "content_status", "identifier": "identifier",
            "language": "language", "version": "version", "created": "created", "modified": "modified",
            "lastPrinted": "last_printed",
        }
        properties: Dict[str, Any] = {}
        for element in root:
            name = etree.QName(element).localname
            if name in field_names and element.text:
                value: Any = element.text
                if name == "revision":
                    try:
                        value = int(value)
                    except ValueError:
                        pass
                properties[field_names[name]] = value
        return properties

    @classmethod
    def _package_application_properties(cls, package: ZipFile) -> Dict[str, str]:
        root = cls._package_xml(package, "docProps/app.xml")
        if root is None:
            return {}
        fields = {"Application": "application", "Template": "template", "AppVersion": "app_version"}
        return {fields[etree.QName(element).localname]: element.text for element in root if etree.QName(element).localname in fields and element.text}

    @classmethod
    def _package_custom_properties(cls, package: ZipFile) -> Dict[str, Any]:
        root = cls._package_xml(package, "docProps/custom.xml")
        if root is None:
            return {}
        properties: Dict[str, Any] = {}
        for property_element in root:
            name = property_element.get("name")
            if not name or not len(property_element):
                continue
            value_element = property_element[0]
            value: Any = value_element.text or ""
            type_name = etree.QName(value_element).localname
            if type_name == "bool":
                value = value.casefold() == "true"
            elif type_name in {"i1", "i2", "i4", "i8", "int", "ui1", "ui2", "ui4", "ui8"}:
                try:
                    value = int(value)
                except ValueError:
                    pass
            elif type_name in {"r4", "r8", "decimal"}:
                try:
                    value = float(value)
                except ValueError:
                    pass
            properties[name] = value
        return properties

    @classmethod
    def _package_protection(cls, package: ZipFile) -> Dict[str, Any]:
        root = cls._package_xml(package, "word/settings.xml")
        if root is None:
            return {"enabled": False}
        protection = root.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}documentProtection")
        if protection is None:
            return {"enabled": False}
        values = {etree.QName(name).localname: value for name, value in protection.attrib.items()}
        return {
            "enabled": values.get("enforcement", "1") not in {"0", "false", "False"},
            **({"edit": values["edit"]} if "edit" in values else {}),
            **({"enforcement": values["enforcement"].casefold() not in {"0", "false"}} if "enforcement" in values else {}),
            **({"formatting": values["formatting"].casefold() not in {"0", "false"}} if "formatting" in values else {}),
        }

    @staticmethod
    def _renderer_path() -> Optional[str]:
        return shutil.which("soffice")

    @staticmethod
    def _renderer_version(renderer: str) -> Optional[str]:
        try:
            completed = subprocess.run([renderer, "--version"], capture_output=True, text=True, timeout=10)
            version = (completed.stdout or completed.stderr).strip()
            return version or None
        except (OSError, subprocess.SubprocessError):
            return None

    def query(
        self,
        path: str,
        query: Dict[str, Any],
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Apply a deliberately closed, document-native filter grammar."""
        unknown_fields = sorted(set(query) - QUERY_FIELDS)
        if unknown_fields:
            return {"ok": False, "error_code": "DOCX_QUERY_INVALID", "message": f"Unsupported query field: {unknown_fields[0]}"}
        if query.get("kind") not in {None, "heading", "paragraph", "list_item", "table"}:
            return {"ok": False, "error_code": "DOCX_QUERY_INVALID", "message": "Unsupported kind filter."}
        if "heading_path" in query and not isinstance(query["heading_path"], list):
            return {"ok": False, "error_code": "DOCX_QUERY_INVALID", "message": "heading_path must be an array."}
        result = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not result["ok"]:
            return result
        applied_filters = {key: value for key, value in query.items() if key != "limit"}
        rows = []
        document = Document(self._resolve_path(path, root_dir, allow_roots, deny_roots))
        block_sections = self._block_sections(document, result["blocks"])
        for ordinal, block in enumerate(result["blocks"], start=1):
            locator = block["locator"]
            if query.get("kind") and block["kind"] != query["kind"]:
                continue
            if query.get("heading_path") is not None and locator["heading_path"] != query["heading_path"]:
                continue
            if query.get("style") and block.get("style") != query["style"]:
                continue
            if query.get("table_id") and block.get("table_id") != query["table_id"]:
                continue
            if query.get("section") is not None and block_sections[ordinal] != query["section"]:
                continue
            text = self._block_text(block)
            if query.get("text") and str(query["text"]).casefold() not in text.casefold():
                continue
            rows.append({"block_id": block["block_id"], "kind": block["kind"], "text": text, "citation": locator})
        bounded_limit = min(max(int(query.get("limit", 50)), 1), RETRIEVAL_RESULT_LIMIT)
        return {
            "ok": True,
            "file": path,
            "applied_filters": applied_filters,
            "results": rows[:bounded_limit],
            "total_results": len(rows),
            "is_truncated": len(rows) > bounded_limit,
            "limit": bounded_limit,
            "fidelity_warnings": result["fidelity_warnings"],
        }

    def structure(
        self,
        path: str,
        root_dir: Optional[str] = None,
        allow_roots: Optional[List[str]] = None,
        deny_roots: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Describe the hierarchy already represented by the ordered block model."""
        result = self.extract_blocks(path, root_dir, allow_roots, deny_roots)
        if not result["ok"]:
            return result
        document = Document(self._resolve_path(path, root_dir, allow_roots, deny_roots))
        blocks = result["blocks"]
        headings = []
        for position, block in enumerate(blocks):
            if block["kind"] != "heading":
                continue
            end = len(blocks)
            for later_position, later in enumerate(blocks[position + 1:], start=position + 2):
                if later["kind"] == "heading" and later["level"] <= block["level"]:
                    end = later_position - 1
                    break
            headings.append({
                "block_id": block["block_id"], "text": block["text"], "level": block["level"],
                "heading_path": block["locator"]["heading_path"],
                "block_range": {"start": position + 1, "end": end},
            })
        return {
            "ok": True,
            "file": path,
            "headings": headings,
            "sections": self._section_ranges(document, blocks),
            "lists": [{"block_id": block["block_id"], "level": block.get("list_level", 1), "text": block["text"]} for block in blocks if block["kind"] == "list_item"],
            "tables": [{"table_id": block["table_id"], "source_index": block["locator"]["source_index"], "heading_path": block["locator"]["heading_path"]} for block in blocks if block["kind"] == "table"],
            "bookmarks": self._bookmarks(document, blocks),
            "headers_footers": {
                "has_headers": any(section.header.paragraphs and any(p.text for p in section.header.paragraphs) for section in document.sections),
                "has_footers": any(section.footer.paragraphs and any(p.text for p in section.footer.paragraphs) for section in document.sections),
            },
            "fidelity_warnings": result["fidelity_warnings"],
        }

    @staticmethod
    def _log_audit(path: str, started_at: float, status: str = "success", **details: Any) -> None:
        logger.info(
            "[DocxAudit] %s",
            {
                "file": path.rsplit("/", maxsplit=1)[-1],
                "elapsed_ms": round((time.monotonic() - started_at) * 1000, 2),
                "status": status,
                **details,
            },
        )

    @staticmethod
    def _searchable_items(blocks: List[Dict[str, Any]]):
        for block in blocks:
            if block["kind"] == "table":
                for row in block["rows"]:
                    for cell in row:
                        yield cell["text"], cell["locator"], "table_cells"
            elif block["kind"] == "heading":
                yield block["text"], block["locator"], "headings"
            else:
                yield block["text"], block["locator"], "body"

    @staticmethod
    def _block_text(block: Dict[str, Any]) -> str:
        if block["kind"] == "table":
            return "\n".join(cell["text"] for row in block["rows"] for cell in row)
        return block["text"]

    @staticmethod
    def _snippet(text: str, span: Dict[str, int], radius: int = 80) -> str:
        start = max(span["start"] - radius, 0)
        end = min(span["end"] + radius, len(text))
        return text[start:end]

    @staticmethod
    def _bookmarks(document: Any, blocks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        body_elements = list(document.element.body.iterchildren())
        locator_by_source = {block["locator"]["source_index"]: block["locator"] for block in blocks}
        bookmarks = []
        for bookmark in document.element.body.iter(f"{namespace}bookmarkStart"):
            owner = next((ancestor for ancestor in bookmark.iterancestors() if ancestor in body_elements), None)
            item: Dict[str, Any] = {
                "name": bookmark.get(f"{namespace}name", ""),
                "id": bookmark.get(f"{namespace}id", ""),
            }
            if owner is not None:
                item["citation"] = locator_by_source.get(body_elements.index(owner) + 1)
            bookmarks.append(item)
        return bookmarks

    @staticmethod
    def _section_ranges(document: Any, blocks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Map OOXML paragraph section breaks to the stable body-block coordinates."""
        block_count = len(blocks)
        if not block_count:
            return [{"index": index, "block_range": {"start": 1, "end": 0}} for index, _ in enumerate(document.sections, start=1)]
        source_to_ordinal = {block["locator"]["source_index"]: ordinal for ordinal, block in enumerate(blocks, start=1)}
        section_ends = []
        for source_index, element in enumerate(document.element.body.iterchildren(), start=1):
            if isinstance(element, CT_P) and element.pPr is not None and element.pPr.sectPr is not None:
                section_ends.append(source_to_ordinal[source_index])
        ranges = []
        start = 1
        for index, end in enumerate(section_ends, start=1):
            ranges.append({"index": index, "block_range": {"start": start, "end": end}})
            start = end + 1
        ranges.append({"index": len(ranges) + 1, "block_range": {"start": start, "end": block_count}})
        return ranges

    @classmethod
    def _block_sections(cls, document: Any, blocks: List[Dict[str, Any]]) -> Dict[int, int]:
        result = {}
        for section in cls._section_ranges(document, blocks):
            for ordinal in range(section["block_range"]["start"], section["block_range"]["end"] + 1):
                result[ordinal] = section["index"]
        return result

    @staticmethod
    def _locator(block_id: str, kind: str, source_index: int, heading_path: List[str]) -> Dict[str, Any]:
        return {
            "block_id": block_id,
            "kind": kind,
            "source_index": source_index,
            "heading_path": list(heading_path),
        }

    @staticmethod
    def _markdown(blocks: List[Dict[str, Any]]) -> str:
        lines: List[str] = []
        for block in blocks:
            if block["kind"] == "heading":
                lines.append(f"{'#' * block['level']} {block['text']}")
            elif block["kind"] == "table":
                for row in block["rows"]:
                    lines.append("| " + " | ".join(cell["text"] for cell in row) + " |")
            else:
                lines.append(block["text"])
        return "\n\n".join(lines)

    @staticmethod
    def _paragraph_block(paragraph: Paragraph, index: int) -> Dict[str, Any]:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        heading = re.fullmatch(r"Heading\s+(\d+)", style_name, flags=re.IGNORECASE)
        block_id = f"p-{index:04d}"
        if heading:
            kind = "heading"
        elif style_name.lower().startswith("list"):
            kind = "list_item"
        else:
            kind = "paragraph"

        block: Dict[str, Any] = {
            "block_id": block_id,
            "kind": kind,
            "text": paragraph.text,
            "style": style_name,
        }
        if heading:
            block["level"] = int(heading.group(1))
        if kind == "list_item":
            number_properties = getattr(getattr(paragraph._p, "pPr", None), "numPr", None)
            nesting_level = getattr(number_properties, "ilvl", None)
            block["list_level"] = int(nesting_level.val) + 1 if nesting_level is not None and nesting_level.val is not None else 1
        return block

    @staticmethod
    def _table_block(table: Table, index: int, heading_path: List[str]) -> Dict[str, Any]:
        table_id = f"table-{index:04d}"
        rows = []
        cell_positions: Dict[int, List[tuple[int, int]]] = {}
        for row_index, row in enumerate(table.rows, start=1):
            cells = []
            for column_index, cell in enumerate(row.cells, start=1):
                cell_positions.setdefault(id(cell._tc), []).append((row_index, column_index))
                cells.append(
                    {
                        "text": cell.text,
                        "locator": {
                            **DocxService._locator(table_id, "table_cell", index, heading_path),
                            "row": row_index,
                            "column": column_index,
                        },
                    }
                )
            rows.append(cells)
        merged_cells = []
        for positions in cell_positions.values():
            if len(positions) > 1:
                rows_in_merge, columns_in_merge = zip(*positions)
                merged_cells.append(
                    {
                        "start": {"row": min(rows_in_merge), "column": min(columns_in_merge)},
                        "end": {"row": max(rows_in_merge), "column": max(columns_in_merge)},
                    }
                )
        return {
            "block_id": table_id,
            "table_id": table_id,
            "kind": "table",
            "locator": DocxService._locator(table_id, "table", index, heading_path),
            "rows": rows,
            "merged_cells": merged_cells,
        }


docx_service = DocxService()
